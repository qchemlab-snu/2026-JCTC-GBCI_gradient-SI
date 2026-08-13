"""Compare analytic and finite-difference GBCI gradients for LiX molecules.

It runs LiH with CAS(2e,2o) and LiF/LiCl with CAS(4e,4o), then compares the
analytic z-gradient on the second atom with a central finite-difference
derivative along the bond.  The GBCI kernel is run before requesting the
gradient so that the current cached-intermediate gradient API is used.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

import numpy as np
import pyscf
import pyscf.grad
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pyscf import gbci, gto, lib, scf

lib.num_threads(1)

REPO_ROOT = Path(__file__).resolve().parents[3]
LOCAL_PYSCF = REPO_ROOT / "pyscf"
LOCAL_PYSCF_GRAD = LOCAL_PYSCF / "grad"
repo_root = str(REPO_ROOT)
if repo_root in sys.path:
    sys.path.remove(repo_root)
sys.path.insert(0, repo_root)
if str(LOCAL_PYSCF) not in pyscf.__path__:
    pyscf.__path__.insert(0, str(LOCAL_PYSCF))
if str(LOCAL_PYSCF_GRAD) not in pyscf.grad.__path__:
    pyscf.grad.__path__.insert(0, str(LOCAL_PYSCF_GRAD))

from pyscf.gbci.gbci import GBCI, biorthogonalize


@dataclass(frozen=True)
class MoleculeSpec:
    name: str
    partner: str
    bond_length: float
    basis: str = "ccpvdz"
    charge: int = 0
    spin: int = 0
    unit: str = "Angstrom"
    ncas: int = 5
    nelecas: tuple[int, int] = (2, 2)
    group_a: dict[str, list[int]] = field(default_factory=lambda: {"atom": [0]})


MOLECULES = [
    MoleculeSpec(name="LiH", partner="H", bond_length=1.50, ncas=2, nelecas=(1, 1)),
    MoleculeSpec(name="LiF", partner="F", bond_length=1.50, ncas=4, nelecas=(2, 2)),
    MoleculeSpec(name="LiCl", partner="Cl", bond_length=1.50, ncas=4, nelecas=(2, 2)),
]


def build_molecule(spec: MoleculeSpec, bond_length: float, verbose: int) -> gto.Mole:
    mol = gto.Mole()
    mol.verbose = verbose
    mol.output = None
    mol.atom = [["Li", (0.0, 0.0, 0.0)], [spec.partner, (0.0, 0.0, bond_length)]]
    mol.basis = spec.basis
    mol.charge = spec.charge
    mol.spin = spec.spin
    mol.unit = spec.unit
    mol.build()
    mol.set_common_orig([0.0, 0.0, 0.0])
    return mol


def run_scf(mol: gto.Mole, conv_tol: float) -> scf.hf.RHF:
    mf = scf.RHF(mol)
    mf.conv_tol = conv_tol
    mf.kernel()
    if not mf.converged:
        raise RuntimeError("RHF did not converge")
    return mf


def format_float_array(values: np.ndarray) -> str:
    return np.array2string(
        np.asarray(values, dtype=float),
        precision=12,
        separator=", ",
        suppress_small=False,
    )


def group_metadata(
    index: int,
    po_list: np.ndarray,
    group: list[list[int]] | None,
) -> tuple[str, list[int], list[list[int]]]:
    if group is None:
        occupation = po_list[index].astype(int).tolist()
        return f"pattern {index} occ={occupation}", [int(index)], [occupation]

    members = [int(member) for member in group[index]]
    occupations = [po_list[member].astype(int).tolist() for member in members]
    return f"group {index} patterns={members} occ={occupations}", members, occupations


def format_group_label(index: int, po_list: np.ndarray, group: list[list[int]] | None) -> str:
    label, _, _ = group_metadata(index, po_list, group)
    return label


def pairwise_differences(values: np.ndarray) -> np.ndarray:
    return np.array(
        [values[i] - values[j] for i in range(len(values)) for j in range(i + 1, len(values))],
        dtype=float,
    )


def float_list(values: np.ndarray) -> list[float]:
    return [float(value) for value in np.asarray(values, dtype=float)]


def print_svd_details(
    label: str,
    gbci_method: GBCI,
    mo_list: np.ndarray,
    po_list: np.ndarray,
    group: list[list[int]] | None,
    ov_list: np.ndarray,
) -> list[dict[str, object]]:
    ncore = gbci_method.ncore
    core_list = np.arange(ncore)
    num_group = len(group) if group is not None else len(po_list)
    s1e = gbci_method._scf.get_ovlp(gbci_method.mol)
    rows = []

    print(f"\nSVD singular values for {label}")
    print(f"Number of group entries: {num_group}, ncore: {ncore}")
    for index in range(num_group):
        print(f"  {format_group_label(index, po_list, group)}")

    print("Ordered group-pair SVD details:")
    for i in range(num_group):
        xc_mo_coeff = mo_list[i][:, core_list]
        for j in range(num_group):
            wc_mo_coeff = mo_list[j][:, core_list]
            singular_values, *_ = biorthogonalize(xc_mo_coeff, wc_mo_coeff, s1e)
            singular_values_squared = singular_values**2
            squared_pairwise_diffs = pairwise_differences(singular_values_squared)
            nonzero_squared = singular_values_squared[np.abs(singular_values) > 1e-10]
            squared_product = float(np.prod(nonzero_squared))
            group_i_label, group_i_members, group_i_occupations = group_metadata(i, po_list, group)
            group_j_label, group_j_members, group_j_occupations = group_metadata(j, po_list, group)
            rows.append(
                {
                    "molecule": label,
                    "ncore": int(ncore),
                    "group_i": int(i),
                    "group_j": int(j),
                    "group_i_label": group_i_label,
                    "group_j_label": group_j_label,
                    "group_i_patterns": group_i_members,
                    "group_j_patterns": group_j_members,
                    "group_i_occupations": group_i_occupations,
                    "group_j_occupations": group_j_occupations,
                    "singular_values": float_list(singular_values),
                    "singular_values_squared": float_list(singular_values_squared),
                    "squared_pairwise_diffs_i_lt_j": float_list(squared_pairwise_diffs),
                    "prod_nonzero_squared": squared_product,
                    "ov_list": float(ov_list[i, j]),
                }
            )
            print(
                f"  group pair ({i}, {j}): "
                f"singular_values={format_float_array(singular_values)}, "
                f"squared={format_float_array(singular_values_squared)}, "
                f"squared_pairwise_diffs_i_lt_j={format_float_array(squared_pairwise_diffs)}, "
                f"prod_nonzero_squared={squared_product:.12e}, "
                f"ov_list={float(ov_list[i, j]):.12e}"
            )
    return rows


def save_svd_logs_to_excel(rows: list[dict[str, object]], output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("script", "pyscf/grad/test/compare_sfnoci_grad_li_halides.py"),
        ("row_count", len(rows)),
        ("note", "Each SVD_Logs row is one ordered group pair for one molecule."),
    ]
    for row in summary_rows:
        summary.append(row)
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 80

    ws = workbook.create_sheet("SVD_Logs")
    max_singular_values = max((len(row["singular_values"]) for row in rows), default=0)
    max_pairwise_diffs = max((len(row["squared_pairwise_diffs_i_lt_j"]) for row in rows), default=0)

    headers = [
        "molecule",
        "ncore",
        "group_i",
        "group_j",
        "group_i_label",
        "group_j_label",
        "group_i_patterns",
        "group_j_patterns",
        "group_i_occupations",
        "group_j_occupations",
        "singular_values",
        "singular_values_squared",
        "squared_pairwise_diffs_i_lt_j",
        "prod_nonzero_squared",
        "ov_list",
    ]
    headers.extend(f"singular_value_{idx}" for idx in range(max_singular_values))
    headers.extend(f"singular_value_squared_{idx}" for idx in range(max_singular_values))
    headers.extend(f"squared_diff_{idx}" for idx in range(max_pairwise_diffs))

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        singular_values = row["singular_values"]
        singular_values_squared = row["singular_values_squared"]
        squared_pairwise_diffs = row["squared_pairwise_diffs_i_lt_j"]
        record = [
            row["molecule"],
            row["ncore"],
            row["group_i"],
            row["group_j"],
            row["group_i_label"],
            row["group_j_label"],
            json.dumps(row["group_i_patterns"]),
            json.dumps(row["group_j_patterns"]),
            json.dumps(row["group_i_occupations"]),
            json.dumps(row["group_j_occupations"]),
            json.dumps(singular_values),
            json.dumps(singular_values_squared),
            json.dumps(squared_pairwise_diffs),
            row["prod_nonzero_squared"],
            row["ov_list"],
        ]
        record.extend(singular_values + [None] * (max_singular_values - len(singular_values)))
        record.extend(singular_values_squared + [None] * (max_singular_values - len(singular_values_squared)))
        record.extend(squared_pairwise_diffs + [None] * (max_pairwise_diffs - len(squared_pairwise_diffs)))
        ws.append(record)

    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(headers, start=1):
        width = max(len(header) + 2, 14)
        if header in {
            "group_i_label",
            "group_j_label",
            "group_i_occupations",
            "group_j_occupations",
            "singular_values",
            "singular_values_squared",
            "squared_pairwise_diffs_i_lt_j",
        }:
            width = 44
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def configure_gbci_solver(gbci_method: GBCI, conv_tol: float) -> None:
    """Set options shared by the GBCI object and its CI solver."""
    for solver in (gbci_method, gbci_method.fcisolver):
        if hasattr(solver, "conv_tol"):
            solver.conv_tol = conv_tol
        if hasattr(solver, "nroots"):
            solver.nroots = 1


def first_energy(energy: object) -> float:
    """Normalize scalar and one-root array return values from GBCI."""
    return float(np.asarray(energy, dtype=float).reshape(-1)[0])


def build_gbci_method(
    spec: MoleculeSpec,
    bond_length: float,
    *,
    scf_conv_tol: float,
    gbci_conv_tol: float,
    verbose: int,
) -> GBCI:
    mol = build_molecule(spec, bond_length, verbose)
    mf = run_scf(mol, scf_conv_tol)
    gbci_method = gbci.gbci(
        mf,
        ncas=spec.ncas,
        nelecas=spec.nelecas,
        group_a=spec.group_a,
    )
    gbci_method._fasscf.conv_tol = 1e-14
    gbci_method._fasscf.conv_tol_grad = 1e-12
    configure_gbci_solver(gbci_method, gbci_conv_tol)
    return gbci_method


def run_gbci_energy(
    spec: MoleculeSpec,
    bond_length: float,
    *,
    scf_conv_tol: float,
    gbci_conv_tol: float,
    verbose: int,
) -> float:
    gbci_method = build_gbci_method(
        spec,
        bond_length,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=gbci_conv_tol,
        verbose=verbose,
    )
    energy, _, _ = gbci_method.kernel()
    return first_energy(energy)


def run_gbci_gradient(
    spec: MoleculeSpec,
    *,
    scf_conv_tol: float,
    gbci_conv_tol: float,
    verbose: int,
    show_svd_details: bool,
    svd_log_rows: list[dict[str, object]] | None,
) -> np.ndarray:
    gbci_method = build_gbci_method(
        spec,
        spec.bond_length,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=gbci_conv_tol,
        verbose=verbose,
    )
    gbci_method.fcisolver.conv_tol = gbci_conv_tol
    gbci_method._fasscf.conv_tol = 1e-14
    gbci_method._fasscf.conv_tol_grad = 1e-12
    gbci_method.kernel()

    if show_svd_details:
        intermediates = gbci_method.get_gbci_intermediates()
        rows = print_svd_details(
            spec.name,
            gbci_method,
            intermediates["mo_list"],
            intermediates["po_list"],
            intermediates["group"],
            intermediates["ov_list"],
        )
        if svd_log_rows is not None:
            svd_log_rows.extend(rows)

    gradient = gbci_method.nuc_grad_method().kernel(state=0)
    return np.asarray(gradient, dtype=float)


def run_sfnoci_energy(
    spec: MoleculeSpec,
    bond_length: float,
    *,
    scf_conv_tol: float,
    fci_conv_tol: float,
    verbose: int,
) -> float:
    """Compatibility wrapper for the older finite-difference sweep script."""
    return run_gbci_energy(
        spec,
        bond_length,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=fci_conv_tol,
        verbose=verbose,
    )


def run_sfnoci_gradient(
    spec: MoleculeSpec,
    *,
    scf_conv_tol: float,
    fci_conv_tol: float,
    verbose: int,
    show_svd_details: bool,
    svd_log_rows: list[dict[str, object]] | None,
) -> np.ndarray:
    """Compatibility wrapper for the older finite-difference sweep script."""
    return run_gbci_gradient(
        spec,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=fci_conv_tol,
        verbose=verbose,
        show_svd_details=show_svd_details,
        svd_log_rows=svd_log_rows,
    )


def compare_gradient(
    spec: MoleculeSpec,
    *,
    delta: float,
    scf_conv_tol: float,
    gbci_conv_tol: float,
    verbose: int,
    show_svd_details: bool,
    svd_log_rows: list[dict[str, object]] | None,
) -> dict[str, object]:
    energy_minus = run_gbci_energy(
        spec,
        spec.bond_length - delta,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=gbci_conv_tol,
        verbose=verbose,
    )
    energy_plus = run_gbci_energy(
        spec,
        spec.bond_length + delta,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=gbci_conv_tol,
        verbose=verbose,
    )
    analytic_gradient = run_gbci_gradient(
        spec,
        scf_conv_tol=scf_conv_tol,
        gbci_conv_tol=gbci_conv_tol,
        verbose=verbose,
        show_svd_details=show_svd_details,
        svd_log_rows=svd_log_rows,
    )

    angstrom_to_bohr = 1.0 / lib.param.BOHR
    numerical_z = (energy_plus - energy_minus) / (2.0 * delta * angstrom_to_bohr)
    analytic_z = float(analytic_gradient[1, 2])
    diff = numerical_z - analytic_z

    return {
        "molecule": asdict(spec),
        "delta_angstrom": delta,
        "energy_minus_hartree": energy_minus,
        "energy_plus_hartree": energy_plus,
        "analytic_gradient_hartree_per_bohr": analytic_gradient.tolist(),
        "analytic_z_atom1_hartree_per_bohr": analytic_z,
        "numerical_z_atom1_hartree_per_bohr": float(numerical_z),
        "difference_hartree_per_bohr": float(diff),
        "abs_difference_hartree_per_bohr": float(abs(diff)),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare analytic and numerical GBCI gradients for LiH CAS(2e,2o) "
            "and LiF/LiCl CAS(4e,4o)."
        )
    )
    parser.add_argument("--delta", type=float, default=1e-5, help="Finite-difference step in Angstrom.")
    parser.add_argument("--scf-conv-tol", type=float, default=1e-12, help="RHF convergence tolerance.")
    parser.add_argument(
        "--gbci-conv-tol",
        "--fci-conv-tol",
        dest="gbci_conv_tol",
        type=float,
        default=1e-12,
        help="GBCI convergence tolerance (legacy alias: --fci-conv-tol).",
    )
    parser.add_argument("--verbose", type=int, default=4, help="PySCF verbosity level.")
    parser.add_argument(
        "--skip-svd-details",
        action="store_true",
        help="Do not collect, print, or save singular-value details for each group pair.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "outputs" / "li_halide_gbci_gradient_comparison.json",
        help="JSON output path.",
    )
    parser.add_argument(
        "--svd-output",
        type=Path,
        default=Path(__file__).resolve().parent / "outputs" / "li_halide_gbci_svd_logs.xlsx",
        help="Excel output path for singular-value logs.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    results = []
    svd_log_rows = []

    for spec in MOLECULES:
        print(f"\n=== {spec.name} ===")
        result = compare_gradient(
            spec,
            delta=args.delta,
            scf_conv_tol=args.scf_conv_tol,
            gbci_conv_tol=args.gbci_conv_tol,
            verbose=args.verbose,
            show_svd_details=not args.skip_svd_details,
            svd_log_rows=svd_log_rows,
        )
        results.append(result)
        print(f"Analytic gradient:  {result['analytic_z_atom1_hartree_per_bohr']:.12f}")
        print(f"Numerical gradient: {result['numerical_z_atom1_hartree_per_bohr']:.12f}")
        print(f"Diff:               {result['difference_hartree_per_bohr']:.12f}")

    payload = {
        "script": "pyscf/grad/test/compare_sfnoci_grad_li_halides.py",
        "results": results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"\nSaved results to {args.output}")
    if not args.skip_svd_details:
        save_svd_logs_to_excel(svd_log_rows, args.svd_output)
        print(f"Saved SVD logs to {args.svd_output}")


if __name__ == "__main__":
    main()
