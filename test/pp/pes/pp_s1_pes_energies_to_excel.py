from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
from ase import io
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
PP_DIR = SCRIPT_DIR.parent
REPO_ROOT = Path(__file__).resolve().parents[3]
repo_root = str(REPO_ROOT)
if repo_root in sys.path:
    sys.path.remove(repo_root)
sys.path.insert(0, repo_root)

from utils.neb_calculator import HARTREE_TO_EV, build_pyscf_mol_from_ase  # noqa: E402

OUTPUT_ROOT = SCRIPT_DIR / "pp_s1_pes_image_outputs"
OUTPUT_XLSX = SCRIPT_DIR / "pp_s1_pes_energies.xlsx"

BASIS = "ccpvdz"
CHARGE = 0
SPIN = 0
UNIT = "Angstrom"
VERBOSE = 4

CASCI_SCANNER_KWARGS = {
    "ncas": 2,
    "nelecas": (1, 1),
    "act_list": [38, 39],
    "nroots": 3,
    "target_root": 1,
    "conv_tol": 1e-10,
    "max_cycle": 200,
    "act_base": 1,
}

GBCI_SCANNER_KWARGS = {
    **CASCI_SCANNER_KWARGS,
    "groupA": {
        "mol1": list(range(0, 9)),
        "mol2": list(range(9, 19)),
    },
}

SEGMENTS = [
    {
        "method": "CASCI",
        "segment": "ground_state_geometry_to_planar",
        "start_label": "ground_state_geometry",
        "end_label": "planar",
        "directory": OUTPUT_ROOT / "casci_ground_to_planar",
        "pattern": "pp_s1_pes_casci_ground_to_planar_*.xyz",
        "results_npz": "pp_s1_pes_casci_ground_to_planar_results.npz",
        "energy_cache_npz": "pp_s1_pes_casci_ground_to_planar_single_point_energies.npz",
    },
    {
        "method": "CASCI",
        "segment": "planar_to_TICT",
        "start_label": "planar",
        "end_label": "TICT",
        "directory": PP_DIR / "casci" / "pp_casci_S1_neb_outputs",
        "pattern": "pp_casci_S1_neb_*.xyz",
        "results_npz": "pp_casci_S1_neb_results.npz",
        "energy_cache_npz": "pp_casci_S1_neb_single_point_energies.npz",
    },
    {
        "method": "GBCI",
        "segment": "ground_state_geometry_to_planar",
        "start_label": "ground_state_geometry",
        "end_label": "planar",
        "directory": OUTPUT_ROOT / "gbci_ground_to_planar",
        "pattern": "pp_s1_pes_gbci_ground_to_planar_*.xyz",
        "results_npz": "pp_s1_pes_gbci_ground_to_planar_results.npz",
        "energy_cache_npz": "pp_s1_pes_gbci_ground_to_planar_single_point_energies.npz",
    },
    {
        "method": "GBCI",
        "segment": "planar_to_TICT",
        "start_label": "planar",
        "end_label": "TICT",
        "directory": OUTPUT_ROOT / "gbci_planar_to_tict_trajectory",
        "pattern": "pp_s1_pes_gbci_planar_to_tict_trajectory_*.xyz",
        "results_npz": "pp_s1_pes_gbci_planar_to_tict_trajectory_results.npz",
        "energy_cache_npz": (
            "pp_s1_pes_gbci_planar_to_tict_trajectory_single_point_energies.npz"
        ),
    },
]

ENERGY_KEYS = (
    "energies_ev",
    "gbci_energies_ev",
    "casci_energies_ev",
    "s1_energies_ev",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Collect PP S1 PES image energies into an Excel workbook."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=OUTPUT_XLSX,
        help="Excel workbook path.",
    )
    parser.add_argument(
        "--reference",
        choices=("min", "first"),
        default="min",
        help="Relative-energy zero for each method.",
    )
    parser.add_argument(
        "--allow-missing",
        action="store_true",
        help="Skip missing segments instead of raising an error.",
    )
    parser.add_argument(
        "--no-compute",
        action="store_true",
        help="Only read finite cached/NEB energies; do not run single-point calculations.",
    )
    return parser.parse_args()


def relative_to_repo(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def read_xyz_positions(path: Path) -> np.ndarray:
    lines = path.read_text().splitlines()
    if not lines:
        raise ValueError(f"Empty xyz file: {path}")
    natm = int(lines[0].strip())
    atom_lines = lines[2 : 2 + natm]
    if len(atom_lines) != natm:
        raise ValueError(f"Expected {natm} atoms in {path}, found {len(atom_lines)}")
    positions = []
    for line in atom_lines:
        fields = line.split()
        if len(fields) < 4:
            raise ValueError(f"Invalid xyz atom line in {path}: {line!r}")
        positions.append([float(value) for value in fields[1:4]])
    return np.asarray(positions, dtype=float)


def make_energy_tracker(method: str, target_root: int = 1) -> Any:
    if method == "CASCI":
        from utils.active_space_tracking_optimizer import (  # noqa: WPS433
            CASCI_Active_Root_Tracking_Optimizer,
        )

        kwargs = {**CASCI_SCANNER_KWARGS, "target_root": target_root}
        return CASCI_Active_Root_Tracking_Optimizer(**kwargs)

    if method == "GBCI":
        from utils.active_space_tracking_optimizer import (  # noqa: WPS433
            GBCI_Active_Root_Tracking_Optimizer,
        )

        kwargs = {**GBCI_SCANNER_KWARGS, "target_root": target_root}
        return GBCI_Active_Root_Tracking_Optimizer(**kwargs)

    raise ValueError(f"Unknown method: {method}")


def run_root_energy_single_point(
    tracker: Any,
    mol: Any,
    method: str,
) -> tuple[np.ndarray, int]:
    mf = tracker.make_mf(mol)
    if method == "GBCI":
        mf.damp = 0.4
        mf.level_shift = 0.5
    mf.kernel()
    if not mf.converged:
        print("WARNING: RHF did not converge.")

    cas_obj = tracker.make_cas(mf)
    mo_cas, selected_act, active_diag = tracker.make_cas_mo(cas_obj, mf)

    if method == "GBCI":
        tracker.run_gbci(cas_obj, mo_cas)
    else:
        cas_obj.kernel(mo_cas)

    method_converged = bool(np.all(getattr(cas_obj, "converged", True)))
    tracker.converged = bool(mf.converged) and method_converged
    if not method_converged:
        print(f"WARNING: {method} did not converge.")

    root, root_diag = tracker.choose_root(cas_obj)
    tracker.current_root = root
    root_energies = tracker.get_root_energies(cas_obj)
    energy_hartree = float(root_energies[root])

    ci_list = tracker._as_ci_list(cas_obj)
    if root < len(ci_list):
        tracker.prev_ci_target = np.array(ci_list[root], copy=True)
    tracker.prev_e_target = energy_hartree
    tracker.prev_mol = mol.copy()
    tracker.prev_coords = mol.atom_coords().copy()
    tracker.act_list = np.asarray(selected_act, dtype=int).copy()

    if method == "GBCI":
        ncore = tracker.get_ncore(cas_obj, mf)
        tracker.prev_active_mo = mo_cas[:, ncore : ncore + tracker.ncas].copy()
        tracker.last_mf = mf
        tracker.last_gbci = cas_obj
        tracker.last_mo_gbci = mo_cas
        tracker.last_selected_act = np.asarray(selected_act, dtype=int).copy()
        tracker.last_active_diag = active_diag
        tracker.last_root_diag = root_diag
    else:
        tracker.prev_active_mo = tracker.get_active_mo_for_tracking(
            cas_obj,
            mo_cas,
        ).copy()

    print(
        f"  tracked root {root}: {energy_hartree:.12f} Eh "
        f"({energy_hartree * HARTREE_TO_EV:.8f} eV)"
    )
    tracker.step += 1
    return np.asarray(root_energies, dtype=float), int(root)


def run_energy_only_single_point(tracker: Any, mol: Any, method: str) -> float:
    root_energies, root = run_root_energy_single_point(tracker, mol, method)
    return float(root_energies[root])


def load_energy_values(segment: dict[str, Any], n_images: int) -> np.ndarray | None:
    directory = Path(segment["directory"])
    candidate_paths = [
        directory / segment["energy_cache_npz"],
        directory / segment["results_npz"],
    ]

    for result_path in candidate_paths:
        if not result_path.exists():
            continue
        with np.load(result_path) as data:
            for key in ENERGY_KEYS:
                if key not in data:
                    continue
                energy_values = np.asarray(data[key], dtype=float)
                if len(energy_values) != n_images:
                    continue
                if np.all(np.isfinite(energy_values)):
                    return energy_values
                return None
    return None


def calculate_segment_energies(
    segment: dict[str, Any],
    image_paths: list[Path],
) -> np.ndarray:
    tracker = make_energy_tracker(segment["method"])
    energies_ev = []

    print(f"Computing {segment['method']} S1 single-point energies: {segment['segment']}")
    for idx, image_path in enumerate(image_paths):
        atoms = io.read(str(image_path))
        mol = build_pyscf_mol_from_ase(
            atoms,
            basis=BASIS,
            charge=CHARGE,
            spin=SPIN,
            unit=UNIT,
            verbose=VERBOSE,
        )
        print(f"  image {idx:03d}: {image_path}")
        energy_hartree = run_energy_only_single_point(
            tracker=tracker,
            mol=mol,
            method=segment["method"],
        )
        energy_ev = float(energy_hartree) * HARTREE_TO_EV
        energies_ev.append(energy_ev)

    return np.asarray(energies_ev, dtype=float)


def save_energy_cache(segment: dict[str, Any], energies_ev: np.ndarray) -> None:
    directory = Path(segment["directory"])
    directory.mkdir(parents=True, exist_ok=True)
    cache_path = directory / segment["energy_cache_npz"]
    np.savez(
        cache_path,
        image_indices=np.arange(len(energies_ev), dtype=int),
        energies_ev=np.asarray(energies_ev, dtype=float),
        energies_hartree=np.asarray(energies_ev, dtype=float) / HARTREE_TO_EV,
    )
    print(f"Saved energy cache: {cache_path}")


def read_segment(
    segment: dict[str, Any],
    allow_missing: bool,
    compute_missing: bool,
) -> list[dict[str, Any]]:
    directory = Path(segment["directory"])
    image_paths = sorted(directory.glob(segment["pattern"]))

    if not image_paths:
        if allow_missing:
            print(f"Skipping missing segment: {segment['method']} {segment['segment']}")
            return []
        raise FileNotFoundError(
            f"Missing xyz images for {segment['method']} {segment['segment']} in {directory}. "
            "Run pp_s1_pes_build_images.py first for new segments."
        )

    energy_values = load_energy_values(segment, n_images=len(image_paths))
    if energy_values is None:
        if not compute_missing:
            raise ValueError(
                f"No finite S1 energies found for {segment['method']} {segment['segment']} "
                f"in {directory}. Rerun without --no-compute to calculate them."
            )
        energy_values = calculate_segment_energies(segment, image_paths)
        save_energy_cache(segment, energy_values)

    if len(image_paths) != len(energy_values) or not np.all(np.isfinite(energy_values)):
        raise ValueError(
            f"Could not obtain finite S1 energies for {segment['method']} "
            f"{segment['segment']}."
        )

    records = []
    for idx, (image_path, energy_ev) in enumerate(zip(image_paths, energy_values)):
        point_label = ""
        if idx == 0:
            point_label = segment["start_label"]
        if idx == len(image_paths) - 1:
            point_label = segment["end_label"]

        records.append(
            {
                "method": segment["method"],
                "segment": segment["segment"],
                "segment_image_index": idx,
                "point_label": point_label,
                "source_xyz": relative_to_repo(image_path),
                "s1_energy_ev": float(energy_ev),
                "positions": read_xyz_positions(image_path),
            }
        )
    return records


def combine_method_records(records: list[dict[str, Any]], reference: str) -> list[dict[str, Any]]:
    combined = []
    for record in records:
        if combined and record["segment_image_index"] == 0:
            continue
        combined.append(record)

    if not combined:
        return []

    distances = [0.0]
    for prev, curr in zip(combined, combined[1:]):
        displacement = curr["positions"] - prev["positions"]
        distances.append(float(np.sqrt(np.mean(displacement**2))))

    path_lengths = np.cumsum(np.asarray(distances, dtype=float))
    total_path = float(path_lengths[-1])
    reaction_coordinate = (
        np.zeros_like(path_lengths) if total_path == 0.0 else path_lengths / total_path
    )

    energies = np.asarray([record["s1_energy_ev"] for record in combined], dtype=float)
    zero = float(np.nanmin(energies)) if reference == "min" else float(energies[0])

    for idx, record in enumerate(combined):
        record.pop("positions")
        record["path_image_index"] = idx
        record["path_length_angstrom"] = float(path_lengths[idx])
        record["reaction_coordinate"] = float(reaction_coordinate[idx])
        record["relative_energy_ev"] = float(record["s1_energy_ev"] - zero)
        record["is_key_point"] = bool(record["point_label"])
    return combined


def write_sheet(ws: Any, records: list[dict[str, Any]]) -> None:
    headers = [
        "method",
        "path_image_index",
        "segment",
        "segment_image_index",
        "point_label",
        "is_key_point",
        "reaction_coordinate",
        "path_length_angstrom",
        "s1_energy_ev",
        "relative_energy_ev",
        "source_xyz",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for record in records:
        ws.append([record.get(header) for header in headers])

    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(headers, start=1):
        width = max(len(header) + 2, 14)
        if header in {"segment", "point_label", "source_xyz"}:
            width = 34 if header != "source_xyz" else 70
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def save_workbook(method_records: dict[str, list[dict[str, Any]]], output_path: Path, reference: str) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("created_at", datetime.now(timezone.utc).isoformat()),
        ("script", Path(__file__).name),
        ("molecule", "pp"),
        ("state", "S1"),
        ("relative_energy_reference", reference),
        ("note", "reaction_coordinate is normalized cumulative RMS displacement."),
    ]
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 80

    all_records = []
    for method in ("CASCI", "GBCI"):
        records = method_records.get(method, [])
        ws = workbook.create_sheet(f"{method}_PES")
        write_sheet(ws, records)
        all_records.extend(records)

    all_ws = workbook.create_sheet("All_Points")
    write_sheet(all_ws, all_records)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    args = parse_args()
    method_records = {"CASCI": [], "GBCI": []}

    for segment in SEGMENTS:
        segment_records = read_segment(
            segment,
            allow_missing=args.allow_missing,
            compute_missing=not args.no_compute,
        )
        method_records[segment["method"]].extend(segment_records)

    for method, records in method_records.items():
        method_records[method] = combine_method_records(records, reference=args.reference)

    save_workbook(method_records, args.output.expanduser().resolve(), args.reference)

    print("===== PP S1 PES Excel summary =====")
    for method in ("CASCI", "GBCI"):
        print(f"{method:5s} points : {len(method_records[method])}")
    print(f"Reference   : {args.reference}")
    print(f"Output XLSX : {args.output.expanduser().resolve()}")


if __name__ == "__main__":
    main()
