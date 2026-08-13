from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
from ase import io
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
PP_DIR = SCRIPT_DIR.parent
REPO_ROOT = Path(__file__).resolve().parents[3]
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from pp_s1_pes_energies_to_excel import (  # noqa: E402
    BASIS,
    CHARGE,
    HARTREE_TO_EV,
    SPIN,
    UNIT,
    VERBOSE,
    build_pyscf_mol_from_ase,
    make_energy_tracker,
    read_xyz_positions,
    relative_to_repo,
    run_root_energy_single_point,
)


OUTPUT_ROOT = SCRIPT_DIR / "pp_s1_pes_image_outputs"
OUTPUT_XLSX = SCRIPT_DIR / "pp_s0_s1_pes_energies.xlsx"
METHODS = ("CASCI", "GBCI")

SEGMENTS = [
    {
        "method": "CASCI",
        "segment": "ground_state_geometry_to_planar",
        "start_label": "ground_state_geometry",
        "end_label": "planar",
        "directory": OUTPUT_ROOT / "casci_ground_to_planar",
        "pattern": "pp_s1_pes_casci_ground_to_planar_*.xyz",
        "energy_cache_npz": (
            "pp_s0_s1_pes_casci_ground_to_planar_single_point_energies.npz"
        ),
        "s1_source_npz": (
            "pp_s1_pes_casci_ground_to_planar_single_point_energies.npz"
        ),
    },
    {
        "method": "CASCI",
        "segment": "planar_to_TICT",
        "start_label": "planar",
        "end_label": "TICT",
        "directory": PP_DIR / "casci" / "pp_casci_S1_neb_outputs",
        "pattern": "pp_casci_S1_neb_*.xyz",
        "energy_cache_npz": "pp_casci_S0_S1_neb_single_point_energies.npz",
        "s1_source_npz": "pp_casci_S1_neb_results.npz",
    },
    {
        "method": "GBCI",
        "segment": "ground_state_geometry_to_planar",
        "start_label": "ground_state_geometry",
        "end_label": "planar",
        "directory": OUTPUT_ROOT / "gbci_ground_to_planar",
        "pattern": "pp_s1_pes_gbci_ground_to_planar_*.xyz",
        "energy_cache_npz": (
            "pp_s0_s1_pes_gbci_ground_to_planar_single_point_energies.npz"
        ),
        "s1_source_npz": (
            "pp_s1_pes_gbci_ground_to_planar_single_point_energies.npz"
        ),
    },
    {
        "method": "GBCI",
        "segment": "planar_to_TICT",
        "start_label": "planar",
        "end_label": "TICT",
        "directory": PP_DIR / "gbci" / "pp_gbci_S1_planar_tict_neb_outputs",
        "pattern": "pp_gbci_S1_planar_tict_neb_*.xyz",
        "energy_cache_npz": (
            "pp_gbci_S0_S1_planar_tict_neb_single_point_energies.npz"
        ),
        "s1_source_npz": "pp_gbci_S1_planar_tict_neb_results.npz",
    },
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Calculate PP S0 energies and combine them with existing S1 PES energies."
        )
    )
    parser.add_argument("--output", type=Path, default=OUTPUT_XLSX)
    parser.add_argument(
        "--no-compute",
        action="store_true",
        help="Only read complete S0/S1 caches; do not run single-point calculations.",
    )
    return parser.parse_args()


def load_state_pair_cache(
    segment: dict[str, Any],
    n_images: int,
) -> tuple[np.ndarray, np.ndarray] | None:
    cache_path = Path(segment["directory"]) / segment["energy_cache_npz"]
    if not cache_path.exists():
        return None

    with np.load(cache_path) as data:
        if "s0_energies_ev" not in data or "s1_energies_ev" not in data:
            return None
        s0_energies = np.asarray(data["s0_energies_ev"], dtype=float)
        s1_energies = np.asarray(data["s1_energies_ev"], dtype=float)

    if len(s0_energies) != n_images or len(s1_energies) != n_images:
        return None
    if not np.all(np.isfinite(s0_energies)) or not np.all(np.isfinite(s1_energies)):
        return None
    return s0_energies, s1_energies


def calculate_state_pair_energies(
    segment: dict[str, Any],
    image_paths: list[Path],
) -> tuple[np.ndarray, np.ndarray]:
    tracker = make_energy_tracker(segment["method"], target_root=0)
    s0_energies_ev = []

    s1_source_path = Path(segment["directory"]) / segment["s1_source_npz"]
    if not s1_source_path.exists():
        raise FileNotFoundError(
            f"Missing existing S1 energy source for {segment['method']} "
            f"{segment['segment']}: {s1_source_path}"
        )
    with np.load(s1_source_path) as data:
        if "energies_ev" not in data:
            raise KeyError(f"{s1_source_path} does not contain 'energies_ev'.")
        s1_energies_ev = np.asarray(data["energies_ev"], dtype=float)
    if len(s1_energies_ev) != len(image_paths) or not np.all(
        np.isfinite(s1_energies_ev)
    ):
        raise ValueError(
            f"Existing S1 energies do not match the image set for "
            f"{segment['method']} {segment['segment']}: {s1_source_path}"
        )
    print(f"Using existing S1 energies: {s1_source_path}")

    print(
        f"Computing {segment['method']} S0 single-point energies: "
        f"{segment['segment']}"
    )
    for idx, image_path in enumerate(image_paths):
        atoms = io.read(str(image_path))
        mol = build_pyscf_mol_from_ase(
            atoms,
            basis=BASIS,
            charge=CHARGE,
            spin=SPIN,
            unit=UNIT,
            verbose=VERBOSE,
        )
        print(f"  image {idx:03d}: {image_path}")
        root_energies, tracked_root = run_root_energy_single_point(
            tracker=tracker,
            mol=mol,
            method=segment["method"],
        )
        if len(root_energies) < 1:
            raise ValueError(
                f"{segment['method']} returned no root energy for {image_path}."
            )
        if tracked_root < 0 or tracked_root >= len(root_energies):
            raise ValueError(
                f"{segment['method']} selected invalid S0 root {tracked_root} "
                f"for {image_path}."
            )
        s0_energies_ev.append(float(root_energies[tracked_root]) * HARTREE_TO_EV)

    return (
        np.asarray(s0_energies_ev, dtype=float),
        s1_energies_ev,
    )


def save_state_pair_cache(
    segment: dict[str, Any],
    s0_energies_ev: np.ndarray,
    s1_energies_ev: np.ndarray,
) -> None:
    cache_path = Path(segment["directory"]) / segment["energy_cache_npz"]
    np.savez(
        cache_path,
        image_indices=np.arange(len(s0_energies_ev), dtype=int),
        s0_energies_ev=s0_energies_ev,
        s1_energies_ev=s1_energies_ev,
        s0_energies_hartree=s0_energies_ev / HARTREE_TO_EV,
        s1_energies_hartree=s1_energies_ev / HARTREE_TO_EV,
        s1_source_npz=np.asarray(
            relative_to_repo(
                Path(segment["directory"]) / segment["s1_source_npz"]
            )
        ),
    )
    print(f"Saved S0/S1 energy cache: {cache_path}")


def read_segment(
    segment: dict[str, Any],
    compute_missing: bool,
) -> list[dict[str, Any]]:
    directory = Path(segment["directory"])
    image_paths = sorted(directory.glob(segment["pattern"]))
    if not image_paths:
        raise FileNotFoundError(
            f"Missing xyz images for {segment['method']} {segment['segment']} "
            f"in {directory}."
        )

    energies = load_state_pair_cache(segment, len(image_paths))
    if energies is None:
        if not compute_missing:
            raise ValueError(
                f"No complete S0/S1 cache for {segment['method']} "
                f"{segment['segment']} in {directory}."
            )
        energies = calculate_state_pair_energies(segment, image_paths)
        save_state_pair_cache(segment, *energies)

    s0_energies_ev, s1_energies_ev = energies
    records = []
    for idx, image_path in enumerate(image_paths):
        point_label = ""
        if idx == 0:
            point_label = segment["start_label"]
        if idx == len(image_paths) - 1:
            point_label = segment["end_label"]
        records.append(
            {
                "method": segment["method"],
                "segment": segment["segment"],
                "segment_image_index": idx,
                "point_label": point_label,
                "source_xyz": relative_to_repo(image_path),
                "s0_energy_ev": float(s0_energies_ev[idx]),
                "s1_energy_ev": float(s1_energies_ev[idx]),
                "positions": read_xyz_positions(image_path),
            }
        )
    return records


def combine_method_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    combined = []
    for record in records:
        if combined and record["segment_image_index"] == 0:
            continue
        combined.append(record)
    if not combined:
        return []

    distances = [0.0]
    for previous, current in zip(combined, combined[1:]):
        displacement = current["positions"] - previous["positions"]
        distances.append(float(np.sqrt(np.mean(displacement**2))))

    path_lengths = np.cumsum(np.asarray(distances, dtype=float))
    total_path = float(path_lengths[-1])
    reaction_coordinate = (
        np.zeros_like(path_lengths) if total_path == 0.0 else path_lengths / total_path
    )
    reference_energy_ev = float(combined[0]["s0_energy_ev"])

    for idx, record in enumerate(combined):
        record.pop("positions")
        record["path_image_index"] = idx
        record["path_length_angstrom"] = float(path_lengths[idx])
        record["reaction_coordinate"] = float(reaction_coordinate[idx])
        record["reference_energy_ev"] = reference_energy_ev
        record["s0_relative_energy_ev"] = (
            float(record["s0_energy_ev"]) - reference_energy_ev
        )
        record["s1_relative_energy_ev"] = (
            float(record["s1_energy_ev"]) - reference_energy_ev
        )
        record["is_key_point"] = bool(record["point_label"])
    return combined


def write_sheet(ws: Any, records: list[dict[str, Any]]) -> None:
    headers = [
        "method",
        "path_image_index",
        "segment",
        "segment_image_index",
        "point_label",
        "is_key_point",
        "reaction_coordinate",
        "path_length_angstrom",
        "s0_energy_ev",
        "s1_energy_ev",
        "reference_energy_ev",
        "s0_relative_energy_ev",
        "s1_relative_energy_ev",
        "source_xyz",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for record in records:
        ws.append([record.get(header) for header in headers])

    ws.freeze_panes = "A2"
    for column_index, header in enumerate(headers, start=1):
        width = max(len(header) + 2, 14)
        if header in {"segment", "point_label"}:
            width = 34
        elif header == "source_xyz":
            width = 76
        ws.column_dimensions[get_column_letter(column_index)].width = width


def save_workbook(
    method_records: dict[str, list[dict[str, Any]]],
    output_path: Path,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("created_at", datetime.now(timezone.utc).isoformat()),
        ("script", Path(__file__).name),
        ("molecule", "pp"),
        ("states", "S0 and S1"),
        ("relative_energy_reference", "method-specific S0 energy at S0 geometry"),
        ("reaction_coordinate", "normalized cumulative RMS displacement"),
        (
            "GBCI planar_to_TICT source",
            relative_to_repo(
                PP_DIR / "gbci" / "pp_gbci_S1_planar_tict_neb_outputs"
            ),
        ),
    ]
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 90

    all_records = []
    for method in METHODS:
        records = method_records[method]
        write_sheet(workbook.create_sheet(f"{method}_PES"), records)
        all_records.extend(records)
    write_sheet(workbook.create_sheet("All_Points"), all_records)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    args = parse_args()
    method_records = {method: [] for method in METHODS}
    for segment in SEGMENTS:
        method_records[segment["method"]].extend(
            read_segment(segment, compute_missing=not args.no_compute)
        )
    for method in METHODS:
        method_records[method] = combine_method_records(method_records[method])

    output_path = args.output.expanduser().resolve()
    save_workbook(method_records, output_path)
    print("===== PP S0/S1 PES Excel summary =====")
    for method in METHODS:
        print(f"{method:5s} points : {len(method_records[method])}")
    print("Reference   : method-specific S0 energy at S0 geometry")
    print(f"Output XLSX : {output_path}")


if __name__ == "__main__":
    main()
