from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Any

os.environ.setdefault("MPLCONFIGDIR", "/tmp/matplotlib-codex")

import matplotlib.image as mpimg
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


SCRIPT_DIR = Path(__file__).resolve().parent
INPUT_XLSX = SCRIPT_DIR / "pp_s0_s1_pes_energies.xlsx"
OUTPUT_PREFIX = SCRIPT_DIR / "pp_s1_pes_figure"

PANEL_INFO = {
    "CASCI": {"sheet": "CASCI_PES", "panel": "(b)", "color": "#1f77b4"},
    "GBCI": {"sheet": "GBCI_PES", "panel": "(c)", "color": "#d62728"},
}

# Reference curves drawn on top of a panel. The SA-CASSCF energies are single points
# evaluated on the CASCI NEB geometries -- there is no separate CASSCF NEB path -- so
# the reaction coordinate lines up with the CASCI curve point by point. The "//CASCI"
# label states that energy-method//geometry-source split.
OVERLAY_INFO = {
    "CASCI": {
        "method": "CASSCF",
        "label": "SA-CASSCF//CASCI",
        "sheet": "CASSCF_PES",
        # Orange, not green: green vs the red GBCI curve has an OKLab dE of only
        # 3.9 under deuteranopia, and green vs the blue CASCI curve fails tritan.
        "color": "#ff7f0e",
        "linestyle": "--",
        "linewidth": 1.8,
    },
}

POINT_LABELS = {
    "ground_state_geometry": "(i) S0 geom.",
    "planar": "(ii) Planar",
    "TICT": "(iii) TICT",
}

STRUCTURE_IMAGES = {
    "ground_state_geometry": "S0_casci_S0_GM.png",
    "planar": "S1_planar_LM.png",
    "TICT": "S1_twisted_LM.png",
}

STRUCTURE_IMAGE_ORDER = ("ground_state_geometry", "planar", "TICT")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Draw the PP S1 CASCI/GBCI PES figure from an Excel workbook."
    )
    parser.add_argument("--input", type=Path, default=INPUT_XLSX, help="Input Excel file.")
    parser.add_argument(
        "--output-prefix",
        type=Path,
        default=OUTPUT_PREFIX,
        help="Output path without extension.",
    )
    parser.add_argument(
        "--formats",
        nargs="+",
        default=["png", "pdf"],
        help="Figure formats to save, e.g. png pdf svg.",
    )
    parser.add_argument(
        "--no-smooth",
        action="store_true",
        help="Plot straight line segments instead of a PCHIP-smoothed PES curve.",
    )
    parser.add_argument(
        "--shared-y",
        action=argparse.BooleanOptionalAction,
        default=True,
        help=(
            "Use the same y-axis range for the CASCI and GBCI panels "
            "(default: enabled)."
        ),
    )
    parser.add_argument(
        "--image-dir",
        type=Path,
        default=SCRIPT_DIR,
        help="Directory containing structure PNG files.",
    )
    parser.add_argument(
        "--no-casscf-overlay",
        action="store_true",
        help="Do not overlay the SA-CASSCF reference curve on the CASCI panel.",
    )
    parser.add_argument(
        "--no-structure-images",
        action="store_true",
        help="Do not add the molecular structure images above the PES panels.",
    )
    parser.add_argument("--show", action="store_true", help="Show the figure window.")
    return parser.parse_args()


def load_pes_sheet(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"Sheet {sheet_name!r} was not found in {workbook_path}.")

    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError(f"Sheet {sheet_name!r} does not contain PES data.")

    headers = [str(value) for value in rows[0]]
    records = []
    skipped_nonfinite = 0
    for row in rows[1:]:
        record = dict(zip(headers, row))
        if record.get("reaction_coordinate") is None:
            continue

        try:
            reaction_coordinate = float(record["reaction_coordinate"])
            s1_relative_energy = float(
                record.get("s1_relative_energy_ev", record.get("relative_energy_ev"))
            )
            s0_value = record.get("s0_relative_energy_ev")
            s0_relative_energy = None if s0_value is None else float(s0_value)
        except (TypeError, ValueError):
            skipped_nonfinite += 1
            continue

        if not np.isfinite(reaction_coordinate) or not np.isfinite(s1_relative_energy):
            skipped_nonfinite += 1
            continue
        if s0_relative_energy is not None and not np.isfinite(s0_relative_energy):
            skipped_nonfinite += 1
            continue

        record["_plot_x"] = reaction_coordinate
        record["_plot_s0_y"] = s0_relative_energy
        record["_plot_s1_y"] = s1_relative_energy
        records.append(record)

    if skipped_nonfinite:
        print(
            f"Skipped {skipped_nonfinite} non-finite energy rows in {sheet_name}. "
            "Geometry-only rows from --interpolate-only need S1 single-point energies "
            "before they can be plotted."
        )
    if not records:
        raise ValueError(
            f"Sheet {sheet_name!r} has no finite PES energy values. "
            "Run the required single-point energy calculations for the images, "
            "then regenerate the Excel file before plotting."
        )
    return records


def format_point_label(
    point_label: str,
    s1_relative_energy: float,
) -> str:
    return f"{POINT_LABELS[point_label]}\nS1 {s1_relative_energy:.3f} eV"


def set_panel_ylim(ax: Any, y_values: np.ndarray) -> None:
    y_min = float(np.nanmin(y_values))
    y_max = float(np.nanmax(y_values))
    y_span = max(y_max - y_min, 0.1)
    lower_margin = max(0.03, 0.05 * y_span)
    upper_margin = max(0.08, 0.16 * y_span)
    ax.set_ylim(y_min - lower_margin, y_max + upper_margin)


def read_cropped_structure_image(path: Path) -> np.ndarray:
    image = mpimg.imread(path)
    rgb = image[..., :3]
    nonwhite = np.any(rgb < 0.98, axis=2)
    if not np.any(nonwhite):
        return image

    y_idx, x_idx = np.where(nonwhite)
    pad = 24
    y0 = max(int(y_idx.min()) - pad, 0)
    y1 = min(int(y_idx.max()) + pad + 1, image.shape[0])
    x0 = max(int(x_idx.min()) - pad, 0)
    x1 = min(int(x_idx.max()) + pad + 1, image.shape[1])
    return image[y0:y1, x0:x1]


def add_structure_image_row(axes: list[Any], image_dir: Path) -> None:
    for ax, point_label in zip(axes, STRUCTURE_IMAGE_ORDER):
        image_path = image_dir / STRUCTURE_IMAGES[point_label]
        ax.axis("off")
        ax.set_title(POINT_LABELS[point_label], fontsize=8.5, pad=1.0)
        if not image_path.exists():
            ax.text(
                0.5,
                0.5,
                f"Missing\n{image_path.name}",
                transform=ax.transAxes,
                ha="center",
                va="center",
                fontsize=8.0,
                color="#666666",
            )
            continue

        ax.imshow(read_cropped_structure_image(image_path))


def plot_overlay_curve(
    ax: Any,
    overlay_info: dict[str, Any],
    records: list[dict[str, Any]],
    smooth: bool,
) -> None:
    x = np.asarray([record["_plot_x"] for record in records], dtype=float)
    y = np.asarray([record["_plot_s1_y"] for record in records], dtype=float)
    order = np.argsort(x)
    x = x[order]
    y = y[order]

    if smooth and len(x) >= 3 and np.all(np.diff(x) > 1.0e-10):
        xs = np.linspace(float(x.min()), float(x.max()), 500)
        ys = PchipInterpolator(x, y)(xs)
    else:
        xs, ys = x, y

    ax.plot(
        xs,
        ys,
        color=overlay_info["color"],
        linestyle=overlay_info["linestyle"],
        linewidth=overlay_info["linewidth"],
        label=overlay_info["label"],
        zorder=2,
    )
    ax.scatter(x, y, s=16, color=overlay_info["color"], zorder=2)


def plot_panel(
    ax: Any,
    method: str,
    records: list[dict[str, Any]],
    smooth: bool,
    panel_label: str | None = None,
    overlay_records: list[dict[str, Any]] | None = None,
) -> None:
    info = PANEL_INFO[method]
    color = info["color"]
    x = np.asarray([record["_plot_x"] for record in records], dtype=float)
    s1_y = np.asarray([record["_plot_s1_y"] for record in records], dtype=float)

    order = np.argsort(x)
    x = x[order]
    s1_y = s1_y[order]

    state_curves = [(method, s1_y, color, "-", 2.2)]


    for state, y_values, state_color, linestyle, linewidth in state_curves:
        if smooth and len(x) >= 3 and np.all(np.diff(x) > 1.0e-10):
            xs = np.linspace(float(x.min()), float(x.max()), 500)
            ys = PchipInterpolator(x, y_values)(xs)
            ax.plot(
                xs,
                ys,
                color=state_color,
                linestyle=linestyle,
                linewidth=linewidth,
                label=state,
            )
        else:
            ax.plot(
                x,
                y_values,
                color=state_color,
                linestyle=linestyle,
                linewidth=linewidth,
                label=state,
            )
        ax.scatter(
            x,
            y_values,
            s=24,
            color=state_color,
            edgecolor="white",
            linewidth=0.7,
            zorder=3,
        )

    for record in records:
        label = record.get("point_label")
        if label not in POINT_LABELS:
            continue
        px = float(record["_plot_x"])
        py = float(record["_plot_s1_y"])
        xytext = (0, -24)
        ha = "left"
        va = "top"
        if label == "ground_state_geometry":
            xytext = (8, 14)
            va = "bottom"
        elif label == "planar":
            ha = "center"
        elif label == "TICT":
            xytext = (-8, 18)
            ha = "right"
            va = "bottom"
        ax.scatter(
            [px],
            [py],
            s=58,
            color=color,
            edgecolor="black",
            linewidth=0.8,
            zorder=4,
        )
        ax.annotate(
            format_point_label(label, py),
            xy=(px, py),
            xytext=xytext,
            textcoords="offset points",
            ha=ha,
            va=va,
            fontsize=7.8,
            annotation_clip=False,
        )

    if overlay_records:
        plot_overlay_curve(ax, OVERLAY_INFO[method], overlay_records, smooth)
        ax.legend(fontsize=7.5, frameon=False, loc="best")

    ax.text(
        0.02,
        1.04,
        f"{panel_label or info['panel']} {method}",
        transform=ax.transAxes,
        ha="left",
        va="bottom",
        fontsize=11,
        fontweight="bold",
        clip_on=False,
    )
    ax.set_xlabel("Reaction Coordinate")
    ax.set_xlim(-0.03, 1.03)
    ax.grid(False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)


def main() -> None:
    args = parse_args()
    workbook_path = args.input.expanduser().resolve()
    output_prefix = args.output_prefix.expanduser().resolve()
    image_dir = args.image_dir.expanduser().resolve()

    records_by_method = {
        method: load_pes_sheet(workbook_path, info["sheet"])
        for method, info in PANEL_INFO.items()
    }

    overlay_by_panel: dict[str, list[dict[str, Any]]] = {}
    if not args.no_casscf_overlay:
        for panel_method, overlay_info in OVERLAY_INFO.items():
            try:
                overlay_by_panel[panel_method] = load_pes_sheet(
                    workbook_path, overlay_info["sheet"]
                )
            except (KeyError, ValueError) as error:
                print(
                    f"Skipping the {overlay_info['method']} overlay on the "
                    f"{panel_method} panel: {error}"
                )

    if args.no_structure_images:
        fig, axes = plt.subplots(
            1,
            2,
            figsize=(7.2, 3.25),
            sharey=args.shared_y,
        )
        fig.subplots_adjust(left=0.10, right=0.98, bottom=0.18, top=0.96, wspace=0.35)
    else:
        fig = plt.figure(figsize=(7.2, 4.35))
        grid = fig.add_gridspec(
            nrows=2,
            ncols=6,
            height_ratios=[0.72, 3.25],
            hspace=0.28,
            wspace=0.70,
        )
        structure_axes = [
            fig.add_subplot(grid[0, 0:2]),
            fig.add_subplot(grid[0, 2:4]),
            fig.add_subplot(grid[0, 4:6]),
        ]
        axes = [
            fig.add_subplot(grid[1, 0:3]),
            fig.add_subplot(grid[1, 3:6], sharey=None),
        ]
        if args.shared_y:
            axes[1].sharey(axes[0])
        add_structure_image_row(structure_axes, image_dir)
        fig.text(0.10, 0.985, "(a)", ha="left", va="top", fontsize=11, fontweight="bold")
        fig.subplots_adjust(left=0.10, right=0.98, bottom=0.13, top=0.92)

    panel_labels = ("(a)", "(b)") if args.no_structure_images else (None, None)
    for ax, method, panel_label in zip(axes, ("CASCI", "GBCI"), panel_labels):
        plot_panel(
            ax,
            method,
            records_by_method[method],
            smooth=not args.no_smooth,
            panel_label=panel_label,
            overlay_records=overlay_by_panel.get(method),
        )

    if args.shared_y:
        all_y = np.asarray(
            [
                record["_plot_s1_y"]
                for records in (
                    list(records_by_method.values())
                    + list(overlay_by_panel.values())
                )
                for record in records
            ],
            dtype=float,
        )
        for ax in axes:
            set_panel_ylim(ax, all_y)
    else:
        for ax, method in zip(axes, ("CASCI", "GBCI")):
            panel_records = list(records_by_method[method])
            panel_records.extend(overlay_by_panel.get(method, []))
            y_values = np.asarray(
                [record["_plot_s1_y"] for record in panel_records],
                dtype=float,
            )
            set_panel_ylim(ax, y_values)
    axes[0].set_ylabel("Relative Energy (eV)")

    for fmt in args.formats:
        output_path = output_prefix.with_suffix(f".{fmt.lstrip('.')}")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(output_path, dpi=600, bbox_inches="tight")
        print(f"Saved figure: {output_path}")

    if args.show:
        plt.show()
    plt.close(fig)


if __name__ == "__main__":
    main()
